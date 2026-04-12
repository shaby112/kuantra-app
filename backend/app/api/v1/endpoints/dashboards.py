
import json
import uuid
import csv
from uuid import UUID
from io import BytesIO, StringIO
from fastapi import APIRouter, Depends, HTTPException, status, Body
from fastapi.responses import StreamingResponse, Response
from typing import List, Dict, Any
from openpyxl import Workbook
from app.api.deps import get_current_user
from app.core.database import AsyncSessionLocal
from app.db.models import User, Dashboard
from app.schemas.dashboard import (
    PlanningRequest, PlanningResponse, GenerationRequest, DashboardOut, DashboardGenerationOut,
    DashboardConfig, WidgetConfig, GridPosition, DashboardCreate, DashboardUpdate, WidgetExecutionStatus
)
from app.services.dashboard_agent_service import dashboard_agent_service
from sqlalchemy import select

router = APIRouter()


def _serialize_cell_value(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, default=str)
    return value


def _normalize_widget_rows(data: Any) -> tuple[list[str], list[dict[str, Any]]]:
    if not isinstance(data, list) or not data:
        return [], []

    first_row = data[0]
    if isinstance(first_row, dict):
        headers: list[str] = list(first_row.keys())
        header_set = set(headers)
        normalized_rows: list[dict[str, Any]] = []
        for row in data:
            if not isinstance(row, dict):
                continue
            for key in row.keys():
                if key not in header_set:
                    headers.append(key)
                    header_set.add(key)
            normalized_rows.append(row)
        return headers, normalized_rows

    # Fallback for primitive arrays.
    return ["value"], [{"value": row} for row in data]

@router.post("/demo/novamart", response_model=DashboardGenerationOut)
async def generate_novamart_demo(
    current_user: User = Depends(get_current_user)
):
    """
    Generate a flagship NovaMart demo dashboard with pre-crafted SQL queries
    that run against the actual synced data. Discovers schema names dynamically.
    """
    from app.services.duckdb_manager import duckdb_manager

    # Discover the NovaMart schema — find which conn_ schema has novamart_ tables
    try:
        schema_rows = duckdb_manager.execute(
            "SELECT DISTINCT table_schema FROM information_schema.tables "
            "WHERE table_name LIKE 'novamart_%' "
            "AND table_schema NOT LIKE '%_staging' "
            "LIMIT 1"
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to discover NovaMart schema: {e}")

    if not schema_rows:
        raise HTTPException(
            status_code=404,
            detail="NovaMart data not found. Please sync a NovaMart PostgreSQL connection first."
        )

    schema = schema_rows[0]["table_schema"]
    customers = f'"{schema}"."novamart_customers"'
    orders = f'"{schema}"."novamart_orders"'
    order_items = f'"{schema}"."novamart_order_items"'

    # Pre-crafted SQL queries for each widget
    queries = {
        "total_revenue": f"SELECT ROUND(CAST(SUM(total_amount) AS DOUBLE), 2) AS value FROM {orders} WHERE status != 'cancelled'",
        "total_customers": f"SELECT COUNT(*) AS value FROM {customers}",
        "total_orders": f"SELECT COUNT(*) AS value FROM {orders}",
        "avg_order_value": f"SELECT ROUND(CAST(AVG(total_amount) AS DOUBLE), 2) AS value FROM {orders} WHERE status != 'cancelled'",
        "revenue_trend": (
            f"SELECT strftime(date_trunc('month', order_date), '%b %Y') AS month, "
            f"ROUND(CAST(SUM(total_amount) AS DOUBLE), 2) AS revenue "
            f"FROM {orders} WHERE status != 'cancelled' "
            f"GROUP BY date_trunc('month', order_date) ORDER BY date_trunc('month', order_date)"
        ),
        "orders_by_status": (
            f"SELECT status, COUNT(*) AS count "
            f"FROM {orders} GROUP BY status ORDER BY count DESC"
        ),
        "top_customers": (
            f"SELECT c.first_name || ' ' || c.last_name AS customer, "
            f"ROUND(CAST(SUM(o.total_amount) AS DOUBLE), 2) AS revenue, "
            f"COUNT(o.order_id) AS orders "
            f"FROM {orders} o JOIN {customers} c ON o.customer_id = c.customer_id "
            f"WHERE o.status != 'cancelled' "
            f"GROUP BY 1 ORDER BY 2 DESC LIMIT 10"
        ),
        "revenue_by_payment": (
            f"SELECT payment_method, ROUND(CAST(SUM(total_amount) AS DOUBLE), 2) AS revenue "
            f"FROM {orders} WHERE status != 'cancelled' "
            f"GROUP BY payment_method ORDER BY revenue DESC"
        ),
        "customers_by_state": (
            f"SELECT state, COUNT(*) AS customers "
            f"FROM {customers} "
            f"GROUP BY state ORDER BY customers DESC LIMIT 10"
        ),
        "recent_orders": (
            f"SELECT o.order_id, c.first_name || ' ' || c.last_name AS customer, "
            f"o.order_date::DATE AS date, o.status, "
            f"ROUND(CAST(o.total_amount AS DOUBLE), 2) AS amount, o.payment_method "
            f"FROM {orders} o JOIN {customers} c ON o.customer_id = c.customer_id "
            f"ORDER BY o.order_date DESC LIMIT 15"
        ),
    }

    def _serialize_results(rows: list) -> list:
        """Ensure all values are JSON-serializable."""
        import datetime as _dt
        from decimal import Decimal
        serialized = []
        for row in rows:
            clean = {}
            for k, v in row.items():
                if isinstance(v, (_dt.datetime, _dt.date)):
                    clean[k] = v.isoformat()
                elif isinstance(v, Decimal):
                    clean[k] = float(v)
                else:
                    clean[k] = v
            serialized.append(clean)
        return serialized

    # Execute all queries and build widget configs
    widgets = []
    widget_status_list = []

    widget_defs = [
        {"key": "total_revenue", "title": "Total Revenue", "type": "metric", "gp": {"x": 0, "y": 0, "w": 3, "h": 2},
         "prefix": "$", "valueFormat": "compact", "colors": ["emerald"]},
        {"key": "total_customers", "title": "Total Customers", "type": "metric", "gp": {"x": 3, "y": 0, "w": 3, "h": 2},
         "valueFormat": "number", "colors": ["violet"]},
        {"key": "total_orders", "title": "Total Orders", "type": "metric", "gp": {"x": 6, "y": 0, "w": 3, "h": 2},
         "valueFormat": "number", "colors": ["blue"]},
        {"key": "avg_order_value", "title": "Avg Order Value", "type": "metric", "gp": {"x": 9, "y": 0, "w": 3, "h": 2},
         "prefix": "$", "valueFormat": "number", "colors": ["amber"]},
        {"key": "revenue_trend", "title": "Revenue Trend", "type": "area", "gp": {"x": 0, "y": 2, "w": 8, "h": 4},
         "colors": ["emerald", "violet"]},
        {"key": "orders_by_status", "title": "Orders by Status", "type": "donut", "gp": {"x": 8, "y": 2, "w": 4, "h": 4},
         "colors": ["emerald", "violet", "blue", "amber", "rose"]},
        {"key": "top_customers", "title": "Top 10 Customers", "type": "bar", "gp": {"x": 0, "y": 6, "w": 6, "h": 4},
         "colors": ["violet"]},
        {"key": "revenue_by_payment", "title": "Revenue by Payment Method", "type": "bar", "gp": {"x": 6, "y": 6, "w": 6, "h": 4},
         "colors": ["emerald"]},
        {"key": "recent_orders", "title": "Recent Orders", "type": "table", "gp": {"x": 0, "y": 10, "w": 8, "h": 4},
         "colors": ["violet"]},
        {"key": "customers_by_state", "title": "Customers by State", "type": "bar", "gp": {"x": 8, "y": 10, "w": 4, "h": 4},
         "colors": ["blue"]},
    ]

    for wd in widget_defs:
        sql = queries[wd["key"]]
        widget_id = str(uuid.uuid4())
        data = []
        error = None

        try:
            result = duckdb_manager.execute(sql)
            data = _serialize_results(result) if isinstance(result, list) else []
        except Exception as e:
            error = str(e)

        if data:
            keys = list(data[0].keys())
            if wd["type"] in ("metric", "kpi", "number", "stat"):
                # Single-value widgets: use the first key as both index and category
                index_key = keys[0]
                cats = [keys[0]]
            else:
                index_key = keys[0]
                cats = [k for k in keys if k != index_key]
        else:
            index_key = "name"
            cats = ["value"]

        widget = WidgetConfig(
            id=widget_id,
            type=wd["type"],
            title=wd["title"],
            data=data,
            index=index_key,
            categories=cats if cats else ["value"],
            gridPosition=GridPosition(**wd["gp"]),
            colors=wd.get("colors"),
            valueFormatter=wd.get("valueFormat", "number"),
            sql_query=sql,
        )
        widgets.append(widget)
        widget_status_list.append(
            WidgetExecutionStatus(
                widget_id=widget_id,
                status="error" if error else "ok",
                error=error,
                sql=sql,
            )
        )

    dashboard_config = DashboardConfig(widgets=widgets)

    async with AsyncSessionLocal() as db:
        dashboard = Dashboard(
            user_id=current_user.id,
            title="NovaMart Command Center",
            config=dashboard_config.model_dump()
        )
        db.add(dashboard)
        await db.commit()
        await db.refresh(dashboard)
        return {
            "id": dashboard.id,
            "title": dashboard.title,
            "config": dashboard.config,
            "created_at": dashboard.created_at,
            "updated_at": dashboard.updated_at,
            "share_url": None,
            "widget_status": widget_status_list,
        }


@router.post("/planning", response_model=PlanningResponse)
async def plan_dashboard(
    request: PlanningRequest,
    current_user: User = Depends(get_current_user)
):
    """
    Step 1: Interactive Dashboard Planning
    Returns clarifying questions or a finalized structured plan.
    """
    return await dashboard_agent_service.plan_dashboard(request.query, request.history, current_user.id, request.connection_ids)

@router.post("/generate", response_model=DashboardGenerationOut)
async def generate_dashboard(
    request: GenerationRequest,
    current_user: User = Depends(get_current_user)
):
    """
    Step 2: Generate Dashboard from Plan
    Executes SQL for each widget, generates data, and saves to DB.
    """
    widgets = []
    widget_status: List[WidgetExecutionStatus] = []

    # Smart layout: KPI/number widgets are small (3-wide), charts are larger (6-wide)
    kpi_types = {"number", "metric", "kpi", "stat", "gauge", "sparkline"}

    # Build aggregation maps for all visualizations
    all_aggs = []
    for viz in request.plan.visualizations:
        aggs = {}
        for m_name in viz.metrics:
            metric_conf = next((m for m in request.plan.metrics if m.name == m_name), None)
            if metric_conf:
                aggs[m_name] = metric_conf.aggregation
        all_aggs.append(aggs)

    # Generate ALL widget SQL in a single LLM call (batch)
    all_results = await dashboard_agent_service.generate_all_widget_data_batch(
        request.plan.visualizations,
        all_aggs,
        current_user.id,
        connection_ids=request.connection_ids,
    )

    for i, (viz, result) in enumerate(zip(request.plan.visualizations, all_results)):
        title = f"{', '.join(viz.metrics)}"
        if viz.breakdown:
            title += f" by {viz.breakdown}"

        widget_id = str(uuid.uuid4())

        # Determine widget size based on type
        if viz.type in kpi_types:
            w, h = 3, 2
        elif viz.type == "table":
            w, h = 6, 4
        elif viz.type == "donut":
            w, h = 4, 4
        else:
            w, h = 6, 4

        widget = WidgetConfig(
            id=widget_id,
            type=viz.type,
            title=title,
            data=result.get("data", []),
            index=result.get("index", "name"),
            categories=result.get("categories", ["value"]),
            gridPosition=GridPosition(x=0, y=0, w=w, h=h),
            colors=None,
            valueFormatter=None,
            sql_query=result.get("sql")
        )
        widgets.append(widget)

        widget_error = result.get("error")
        widget_status.append(
            WidgetExecutionStatus(
                widget_id=widget_id,
                status="error" if widget_error else "ok",
                error=widget_error,
                sql=result.get("sql"),
            )
        )

    # Auto-layout: pack widgets into a 12-column grid row by row
    cur_x, cur_y, row_max_h = 0, 0, 0
    for widget in widgets:
        gp = widget.gridPosition
        if cur_x + gp.w > 12:
            cur_x = 0
            cur_y += row_max_h
            row_max_h = 0
        gp.x = cur_x
        gp.y = cur_y
        cur_x += gp.w
        row_max_h = max(row_max_h, gp.h)

    dashboard_config = DashboardConfig(widgets=widgets)
    
    # Save to Database
    async with AsyncSessionLocal() as db:
        dashboard = Dashboard(
            user_id=current_user.id,
            title=request.plan.title,
            config=dashboard_config.model_dump()
        )
        db.add(dashboard)
        await db.commit()
        await db.refresh(dashboard)
        return {
            "id": dashboard.id,
            "title": dashboard.title,
            "config": dashboard.config,
            "created_at": dashboard.created_at,
            "updated_at": dashboard.updated_at,
            "share_url": None,
            "widget_status": widget_status,
        }

@router.post("/", response_model=DashboardOut)
async def create_dashboard(
    request: DashboardCreate,
    current_user: User = Depends(get_current_user)
):
    """Create a new dashboard."""
    async with AsyncSessionLocal() as db:
        dashboard = Dashboard(
            user_id=current_user.id,
            title=request.title,
            config=request.config.model_dump()
        )
        db.add(dashboard)
        await db.commit()
        await db.refresh(dashboard)
        return dashboard

@router.put("/{dashboard_id}", response_model=DashboardOut)
async def update_dashboard(
    dashboard_id: UUID,
    request: DashboardUpdate,
    current_user: User = Depends(get_current_user)
):
    """Update an existing dashboard."""
    async with AsyncSessionLocal() as db:
        result = await db.execute(select(Dashboard).where(Dashboard.id == dashboard_id, Dashboard.user_id == current_user.id))
        dashboard = result.scalar_one_or_none()
        if not dashboard:
            raise HTTPException(status_code=404, detail="Dashboard not found")
        
        if request.title is not None:
            dashboard.title = request.title
        if request.config is not None:
            dashboard.config = request.config.model_dump()
            
        await db.commit()
        await db.refresh(dashboard)
        return dashboard

@router.get("/", response_model=List[DashboardOut])
async def list_dashboards(current_user: User = Depends(get_current_user)):
    """List all dashboards for current user."""
    async with AsyncSessionLocal() as db:
        result = await db.execute(select(Dashboard).where(Dashboard.user_id == current_user.id))
        return result.scalars().all()

@router.get("/{dashboard_id}", response_model=DashboardOut)
async def get_dashboard(dashboard_id: UUID, current_user: User = Depends(get_current_user)):
    """Get a specific dashboard."""
    async with AsyncSessionLocal() as db:
        result = await db.execute(select(Dashboard).where(Dashboard.id == dashboard_id, Dashboard.user_id == current_user.id))
        dashboard = result.scalar_one_or_none()
        if not dashboard:
            raise HTTPException(status_code=404, detail="Dashboard not found")
        return dashboard

@router.delete("/{dashboard_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_dashboard(dashboard_id: UUID, current_user: User = Depends(get_current_user)):
    """Delete a specific dashboard."""
    async with AsyncSessionLocal() as db:
        result = await db.execute(
            select(Dashboard).where(Dashboard.id == dashboard_id, Dashboard.user_id == current_user.id)
        )
        dashboard = result.scalar_one_or_none()
        if not dashboard:
            raise HTTPException(status_code=404, detail="Dashboard not found")
        
        await db.delete(dashboard)
        await db.commit()

@router.get("/{dashboard_id}/export/json")
async def export_dashboard_json(dashboard_id: UUID, current_user: User = Depends(get_current_user)):
    """Export dashboard configuration as JSON."""
    dashboard = await get_dashboard(dashboard_id, current_user)
    # config is already a dict from JSON column
    json_str = json.dumps(dashboard.config, default=str, indent=2)
    return Response(
        content=json_str,
        media_type="application/json",
        headers={"Content-Disposition": f"attachment; filename=dashboard_{dashboard_id}.json"}
    )

@router.get("/{dashboard_id}/export/csv")
async def export_dashboard_csv(dashboard_id: UUID, current_user: User = Depends(get_current_user)):
    """Export dashboard data as a multi-section CSV."""
    dashboard = await get_dashboard(dashboard_id, current_user)
    
    output = StringIO()
    # config is a dict, access 'widgets' list
    widgets = dashboard.config.get("widgets", [])
    
    for i, widget in enumerate(widgets):
        title = widget.get("title", f"Widget {i+1}")
        type_ = widget.get("type", "unknown")
        data = widget.get("data", [])
        
        output.write(f"# Widget: {title} ({type_})\n")
        headers, rows = _normalize_widget_rows(data)
        if not headers:
            output.write("No Data\n")
        else:
            writer = csv.DictWriter(output, fieldnames=headers)
            writer.writeheader()
            for row in rows:
                writer.writerow({h: _serialize_cell_value(row.get(h)) for h in headers})
        output.write("\n\n")
    
    return Response(
        content=output.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=dashboard_{dashboard_id}.csv"}
    )

@router.get("/{dashboard_id}/export/excel")
async def export_dashboard_excel(dashboard_id: UUID, current_user: User = Depends(get_current_user)):
    """Export dashboard data as an Excel workbook."""
    dashboard = await get_dashboard(dashboard_id, current_user)
    widgets = dashboard.config.get("widgets", [])
    
    output = BytesIO()
    workbook = Workbook()
    default_sheet = workbook.active

    if not widgets:
        default_sheet.title = "Empty"
        default_sheet.append(["message"])
        default_sheet.append(["No Widgets"])
    else:
        workbook.remove(default_sheet)
        used_sheet_names: set[str] = set()
        for i, widget in enumerate(widgets):
            title = widget.get("title", f"Widget {i+1}")
            sheet_name = f"{i + 1}_{title}"
            for char in [":", "\\", "/", "?", "*", "[", "]"]:
                sheet_name = sheet_name.replace(char, "")
            sheet_name = sheet_name[:31] or f"widget_{i + 1}"

            # Ensure sheet names are unique after truncation.
            base_name = sheet_name
            suffix = 1
            while sheet_name in used_sheet_names:
                suffix_str = f"_{suffix}"
                sheet_name = f"{base_name[:31-len(suffix_str)]}{suffix_str}"
                suffix += 1
            used_sheet_names.add(sheet_name)

            sheet = workbook.create_sheet(title=sheet_name)
            headers, rows = _normalize_widget_rows(widget.get("data", []))
            if not headers:
                sheet.append(["message"])
                sheet.append(["No Data"])
                continue

            sheet.append(headers)
            for row in rows:
                sheet.append([_serialize_cell_value(row.get(h)) for h in headers])

    workbook.save(output)

    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=dashboard_{dashboard_id}.xlsx"}
    )
